#!/usr/bin/env python3
"""Import cached launch content from the read-only balancing workbook.

Source cells are authoritative. Derived values retain their operands and rule;
unscheduled events and ambiguous capacity labels stay explicitly unconfigured.
"""
from __future__ import annotations

import argparse
import calendar
import hashlib
import json
import math
import re
from pathlib import Path

import openpyxl
try:
    from .import_supplement import import_supplement, write_json_atomic
except ImportError:
    from import_supplement import import_supplement, write_json_atomic


CATALOGS = {
    'Coin Fish': {'prefix': 'CF', 'columns': {
        'name': 'B', 'rarity': 'C', 'level': 'D', 'price': 'E', 'stage_ms': 'F',
        'feed_ms': 'I', 'grace_ms': 'J', 'buy_xp': 'M', 'role': 'AA',
        'description': 'AB', 'badge': 'AD',
        'sale_coins': ['R', 'Q', 'P', 'N'], 'sale_xp': ['U', 'T', 'S', 'O']}},
    'Premium Fish': {'prefix': 'PF', 'columns': {
        'name': 'B', 'level': 'C', 'price': 'D', 'stage_ms': 'E',
        'feed_ms': 'H', 'grace_ms': 'I', 'buy_xp': 'J', 'role': 'U',
        'description': 'AA', 'sale_coins': ['O', 'N', 'M', 'K'],
        'sale_xp': ['R', 'Q', 'P', 'L']}},
    'Limited Edition': {'prefix': 'LE', 'columns': {
        'name': 'B', 'event_copy': 'D', 'level': 'E', 'currency': 'F',
        'price': 'G', 'stage_ms': 'H', 'feed_ms': 'K', 'grace_ms': 'L',
        'buy_xp': 'O', 'description': 'Z',
        'sale_coins': ['R', None, None, 'P'], 'sale_xp': ['S', None, None, 'Q']}},
}
INPUT_CELLS = {
    'starting_coins': 'B6', 'starting_pearls': 'B7',
    'starting_tank_capacity': 'B8', 'starter_fish_count': 'B9',
    'level_2_pearl_grant': 'B11', 'growth_stages_to_adult': 'B19',
    'feed_window_stage_multiplier': 'B20', 'max_free_tank_capacity': 'B63',
    'daily_quest_xp_share': 'B78', 'weekly_quest_xp_share': 'B79',
    'daily_quest_coin_share': 'B80',
}


def norm(value):
    return re.sub(r'[^a-z0-9]+', ' ', str(value or '').lower()).strip()


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def rnd(value):
    return int(math.floor(value + .5))


def camel(value):
    words = re.findall(r'[A-Za-z0-9]+', value)
    return words[0].lower() + ''.join(word[0].upper() + word[1:] for word in words[1:])


def source(sheet, *cells, rule=None):
    result = {'sheet': sheet, 'cells': list(cells)}
    if rule:
        result['derivation'] = rule
    return result


def implementation(reason):
    return {'kind': 'implementation', 'reason': reason}


def required_number(sheet, address):
    value = sheet[address].value
    if not number(value):
        raise ValueError(f'{sheet.title}!{address}: missing or nonnumeric cached result')
    return value


def event_window(text):
    """Read month/day windows literally; never assign dates to '1 week, yearly'."""
    if text == '1 week, yearly':
        return 0, 0, False
    match = re.fullmatch(r'([A-Za-z]+) (\d{1,2})-([A-Za-z]+ )?(\d{1,2})(?: only)?', text)
    if not match:
        raise ValueError(f'Unrecognized event availability: {text!r}')
    months = {name.lower(): i for i, name in enumerate(calendar.month_abbr) if name}
    start_month = months[match[1].lower()]
    end_month = months[match[3].strip().lower()] if match[3] else start_month
    first, last = int(match[2]), int(match[4])
    for month, day in [(start_month, first), (end_month, last)]:
        if not 1 <= day <= calendar.monthrange(2000, month)[1]:
            raise ValueError(f'Invalid event day: {text}')
    return start_month * 100 + first, end_month * 100 + last, True


def import_sheet(workbook, title, schema):
    sheet = workbook[title]
    columns = schema['columns']
    prefix = schema['prefix']
    output = []
    # Materialize values once. ReadOnlyWorksheet cell indexing reparses XML.
    rows = list(sheet.values)
    for row_number, values in enumerate(rows, 1):
        model_id = values[0]
        if not isinstance(model_id, str) or not re.fullmatch(prefix + r'-?\d+', model_id):
            continue
        def value(column):
            index = openpyxl.utils.column_index_from_string(column) - 1
            return values[index] if index < len(values) else None
        def numeric(field):
            result = value(columns[field])
            if not number(result):
                raise ValueError(f'{title}!{columns[field]}{row_number}: missing cached {field}')
            return result
        def origin(field, rule=None):
            return source(title, f'{columns[field]}{row_number}', rule=rule)
        name = str(value(columns['name']))
        currency = 'coins' if prefix == 'CF' else 'pearls' if prefix == 'PF' else str(value('F')).lower()
        if currency not in ('coins', 'pearls', 'gift'):
            raise ValueError(f'{model_id}: unknown currency {currency!r}')
        stage_hours = numeric('stage_ms')
        role = str(value(columns['role'])) if 'role' in columns else 'Collection'
        row = {
            'id': camel(name), 'model_id': model_id, 'name': name,
            'rarity': str(value('C')).lower() if prefix == 'CF' else 'premium' if prefix == 'PF' else 'limited',
            'role': role, 'level': int(numeric('level')), 'currency': currency,
            'price': rnd(numeric('price')), 'buy_xp': rnd(numeric('buy_xp')),
            'stage_ms': rnd(stage_hours * 3600000),
            'feed_ms': rnd(numeric('feed_ms') * 3600000),
            'grace_ms': rnd(numeric('grace_ms') * 3600000),
            'sale_coins': [0], 'sale_xp': [0],
            'badge': str(value('AD')) if prefix == 'CF' else 'Fast Fish' if stage_hours < 1 else 'Overnight' if stage_hours >= 24 else '',
            'description': str(value(columns['description']) or role),
            'length': 58 if prefix == 'PF' else 38,
            'event_start': 0, 'event_end': 0, 'event_configured': prefix != 'LE',
            'annual': False, 'one_time': False, 'non_resellable': False,
            'source_row': row_number, 'source_sheet': title,
        }
        provenance = {
            'id': origin('name', 'ASCII words converted to lower camel case'),
            'model_id': source(title, f'A{row_number}'),
            'name': origin('name'), 'level': origin('level'),
            'price': origin('price'), 'buy_xp': origin('buy_xp'),
            'currency': origin('currency', 'lowercase') if prefix == 'LE' else source(title, f'{columns["price"]}1', rule='currency named by purchase column'),
            'rarity': origin('rarity', 'lowercase') if prefix == 'CF' else source(title, 'A1', rule=f'{title} catalog classification'),
            'role': origin('role') if 'role' in columns else implementation('Default role; Limited Edition has no role column'),
            'description': origin('description') if value(columns['description']) else implementation('Role used when description cell is blank'),
            'badge': origin('badge') if prefix == 'CF' else implementation('Display badge derived from stage hours'),
            'length': implementation('Visual sprite size, not a balancing input'),
            'one_time': implementation('Repeat purchase by default'),
        }
        for field in ['stage_ms', 'feed_ms', 'grace_ms']:
            provenance[field] = origin(field, 'hours × 3,600,000, rounded half up')
        for field in ['sale_coins', 'sale_xp']:
            origins = [implementation('Baby fish have zero sale rewards')]
            for age, column in enumerate(columns[field], 1):
                if column:
                    reward = value(column)
                    if not number(reward):
                        raise ValueError(f'{title}!{column}{row_number}: missing cached sale reward')
                    entry = source(title, f'{column}{row_number}')
                else:
                    # Only LE young/mature rewards lack explicit cells.
                    ratio_address = f'B{21 + age}'
                    ratio = required_number(workbook['Inputs'], ratio_address)
                    adult_column = columns[field][-1]
                    adult = value(adult_column)
                    if not number(adult):
                        raise ValueError(f'{model_id}: missing adult sale reward')
                    reward = row['price'] + (adult - row['price']) * ratio if field == 'sale_coins' and currency == 'coins' else adult * ratio
                    entry = {
                        'sources': [source(title, f'{adult_column}{row_number}', f'{columns["price"]}{row_number}'), source('Inputs', ratio_address)],
                        'derivation': 'round_half_up(price + (adult - price) × stage ratio)' if field == 'sale_coins' and currency == 'coins' else 'round_half_up(adult × stage ratio)',
                    }
                row[field].append(rnd(reward))
                origins.append(entry)
            provenance[field] = origins
        if prefix == 'LE':
            row['event_copy'] = str(value('D'))
            row['event_start'], row['event_end'], row['event_configured'] = event_window(row['event_copy'])
            for field in ['event_copy', 'event_start', 'event_end', 'event_configured']:
                provenance[field] = source(title, f'D{row_number}', rule='literal month/day range; zero means calendar not configured')
            if currency == 'gift':
                row.update(annual=True, non_resellable=True)
                provenance['annual'] = source(title, f'D{row_number}', rule='yearly gift claim')
                provenance['non_resellable'] = source(title, f'Y{row_number}', f'Z{row_number}')
        if row['id'] == 'bubbleEyeGoldfish':
            note = str(value('AA'))
            if 'repeatable' not in note.lower():
                raise ValueError('Review Bubble Eye purchase policy: source notes changed')
            provenance['one_time'] = source(title, f'AA{row_number}', rule='Repeatable purchase; detailed catalog note takes precedence over static summary')
        row['provenance'] = provenance
        output.append(row)
    if not output:
        raise ValueError(f'{title}: no {prefix} catalog rows')
    return output, [norm(value) for value in rows[0]]


def slash_numbers(value):
    parts = str(value).replace(',', '').split('/')
    try:
        result = [int(part.strip()) for part in parts]
    except ValueError as error:
        raise ValueError(f'Invalid slash-separated numeric schedule: {value!r}') from error
    if len(result) != 4 or any(item < 0 for item in result):
        raise ValueError(f'Expected four nonnegative schedule entries: {value!r}')
    return result


def import_tanks(workbook):
    result = {key: [] for key in ['tank_costs', 'tank_tokens', 'tank_friend_assists', 'tank_capacities', 'tank_levels', 'tank_expansion_levels']}
    provenance = {key: [] for key in result}
    sheet = workbook['Tanks & Social']
    beats = {}
    for row in workbook['Unlock Plan'].iter_rows(min_row=4, values_only=True):
        if len(row) < 7 or not number(row[0]) or not 1 <= row[0] <= 40:
            continue
        match = re.fullmatch(r'(Medium|Large) Tank (\d+)', str(row[6]))
        if match:
            beats[(int(match[2]), match[1])] = (int(row[0]), f'G{int(row[0]) + 3}')
    for row_number in range(4, 9):
        tank_id = row_number - 3
        for key, column in [('tank_costs', 'H'), ('tank_tokens', 'G'), ('tank_friend_assists', 'F'), ('tank_capacities', 'E')]:
            text = str(sheet[f'{column}{row_number}'].value).removesuffix(' friends')
            result[key].append(slash_numbers(text))
            provenance[key].append(source(sheet.title, f'{column}{row_number}', rule='slash-separated schedule'))
        level = int(required_number(sheet, f'B{row_number}'))
        result['tank_levels'].append(level)
        provenance['tank_levels'].append(source(sheet.title, f'B{row_number}'))
        expansion = [level]
        origins = [source(sheet.title, f'B{row_number}')]
        for label in ['Medium', 'Large']:
            beat = beats.get((tank_id, label))
            expansion.append(beat[0] if beat else None)
            origins.append(source('Unlock Plan', beat[1], rule=f'Explicit {label} beat; association to slot count is unconfirmed') if beat else implementation(f'No {label} beat within launch levels'))
        expansion.append(None)
        origins.append(implementation('Workbook does not assign an unlock level to the fourth capacity step'))
        result['tank_expansion_levels'].append(expansion)
        provenance['tank_expansion_levels'].append(origins)
    result['tank_expansion_mapping_confirmed'] = False
    return result, provenance


def rectangular_rows(sheet):
    rows = [list(row) for row in sheet.values]
    width = max(map(len, rows), default=0)
    return [row + [None] * (width - len(row)) for row in rows]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--out', type=Path, default=Path('assets/content.json'))
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    species, headers = [], {}
    for title, schema in CATALOGS.items():
        items, header = import_sheet(workbook, title, schema)
        species.extend(items)
        headers[title] = header
    counts = {kind: sum(item['model_id'].startswith(prefix) for item in species) for kind, prefix in [('coin', 'CF'), ('premium', 'PF'), ('limited', 'LE')]}
    if counts != {'coin': 26, 'premium': 10, 'limited': 10}:
        raise ValueError(f'Catalog count changed: {counts}; review launch scope')
    rows = list(workbook['XP Curve'].values)
    header_row = max(range(min(15, len(rows))), key=lambda i: sum('cumulative' in norm(cell) for cell in rows[i]))
    header = [norm(cell) for cell in rows[header_row]]
    cumulative_column = next((i for i, value in enumerate(header) if 'proposed' in value and 'cumulative' in value), None)
    if cumulative_column is None:
        cumulative_column = next(i for i, value in enumerate(header) if 'cumulative' in value)
    level_column = next(i for i, value in enumerate(header) if value == 'level')
    curves, curve_origins = {}, {}
    for row_number, row in enumerate(rows[header_row + 1:], header_row + 2):
        if number(row[level_column]) and number(row[cumulative_column]) and 1 <= row[level_column] <= 70:
            level = int(row[level_column])
            if level in curves:
                raise ValueError(f'Duplicate XP level {level}')
            curves[level] = int(row[cumulative_column])
            curve_origins[level] = source('XP Curve', f'{openpyxl.utils.get_column_letter(cumulative_column + 1)}{row_number}')
    if sorted(curves) != list(range(1, 71)) or [curves[1], curves[2]] != [0, 80]:
        raise ValueError('XP curve must contain consecutive levels 1-70 with cached results')
    inputs = {key: required_number(workbook['Inputs'], cell) for key, cell in INPUT_CELLS.items()}
    names = [name.strip() for name in str(workbook['Unlock Plan']['C4'].value).split(',')]
    by_name = {item['name']: item['id'] for item in species}
    starter_species = [by_name[name] for name in names]
    if len(starter_species) != inputs['starter_fish_count']:
        raise ValueError('Starting fish count disagrees with Unlock Plan C4')
    tanks, tank_origins = import_tanks(workbook)
    result = {
        'schema': 1, 'workbook_sha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
        'counts': counts, 'species': species,
        'levels': [curves[level] for level in range(1, 41)],
        'future_levels': [curves[level] for level in range(41, 71)],
        'inputs': inputs, 'starter_species': starter_species,
        **tanks,
        'provenance': {
            'inputs': {key: source('Inputs', cell) for key, cell in INPUT_CELLS.items()},
            'starter_species': source('Unlock Plan', 'C4', rule='names resolved to catalog IDs'),
            'levels': [curve_origins[level] for level in range(1, 41)],
            **tank_origins,
        },
        'source_conflicts': [{
            'field': 'species.bubbleEyeGoldfish.one_time',
            'chosen': source('Premium Fish', 'AA2'),
            'conflicting_summaries': [source('Unlock Plan', 'F5'), source('Balance Review', 'B6'), source('Inputs', 'F111'), source('XP Curve', 'H5')],
            'resolution': 'Detailed catalog note explicitly says repeatable; static one-time summaries are stale.',
        }],
        'unconfigured': [{
            'field': 'species.anniversaryRainbowfish.event_start/event_end',
            'source': source('Limited Edition', 'D11'),
            'reason': 'The workbook specifies one week yearly, but no calendar dates.',
        }, {
            'field': 'tank_expansion_levels',
            'reason': 'Medium and Large beats are explicit; their capacity mapping and fourth-step unlock are not specified.',
        }],
        'raw_sheets': {sheet.title: rectangular_rows(sheet) for sheet in workbook},
        'headers': headers,
    }
    workbook.close()
    if args.out.exists():
        previous = json.loads(args.out.read_text())
        result['supplement'] = previous.get('supplement', {})
    # Standalone imports must be runnable game content too. Rebuild all mapped
    # quest/mastery fields before replacing the destination, preserving only
    # deferred or custom fields that the supplemental importer does not own.
    result = import_supplement(result)
    write_json_atomic(args.out, result)
    print(json.dumps({'output': str(args.out), 'species': len(species), 'launch_levels': 40, 'future_levels': 30, 'unconfigured': result['unconfigured']}, indent=2))


if __name__ == '__main__':
    main()
