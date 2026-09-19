#!/usr/bin/env python3
"""Independently compare runtime cosmetics against the original workbook cells."""
import argparse
import json
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP, ROUND_FLOOR
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,default=ROOT/'design/aquarium_game_design_v4.xlsx')
    parser.add_argument('--content',type=Path,default=ROOT/'assets/content.json')
    parser.add_argument('--out',type=Path,default=ROOT/'evidence/decor-catalog-audit.json')
    args=parser.parse_args()
    data=json.loads(args.content.read_text())['decorations']
    w=openpyxl.load_workbook(args.workbook,read_only=True,data_only=True)
    sheets={s.title:list(s.values) for s in w};w.close()
    errors=[];checks=0
    def same(name,actual,expected,source):
        nonlocal checks
        checks+=1
        if actual!=expected:errors.append(dict(field=name,actual=actual,expected=expected,source=source))
    items={i['id']:i for i in data['items']}
    same('unique IDs',len(items),120,'Decor Catalog!A5:A124')
    same('row count',len(data['items']),120,'Decor Catalog!A5:A124')
    same('categories',dict(Counter(i['category'] for i in data['items'])),{'Plant':60,'Decoration':60},'Decor Catalog!C5:C124')
    columns={'id':0,'name':1,'category':2,'theme':3,'edition':4,'rarity':5,'level':6,'currency':7,'price':11,'score':13,'width':14,'height':15,'size':16,'layer':17,'event':18,'release_gate':19,'subcategory':21}
    for number,row in enumerate(sheets['Decor Catalog'][4:124],5):
        item=items.get(row[0],{})
        for key,col in columns.items():
            expected=row[col].lower() if key=='currency' else '' if key=='event' and row[col]=='None' else row[col]
            same(f'{row[0]}.{key}',item.get(key),expected,f'Decor Catalog!{openpyxl.utils.get_column_letter(col+1)}{number}')
        xp=max(1,int(row[11])//10+(int(row[11])%10>=5)) if row[7]=='Coins' else int(row[11])*10
        same(f'{row[0]}.first_purchase_xp',item.get('buy_xp'),xp,'18 September price-based decor XP override')
        same(f'{row[0]}.asset',item.get('asset'),'decor/catalog/'+row[0]+'.png','Catalog ID join')
        same(f'{row[0]}.source_row',item.get('source',{}).get('row'),number,'Decor Catalog')
    art_fields=['id','name','silhouette','palette','materials','footprint','placement','animation','motion_class','still','hook','construction']
    for number,row in enumerate(sheets['Decor Art Briefs'][4:124],5):
        art=items.get(row[0],{}).get('art',{})
        for col,key in enumerate(art_fields):same(f'{row[0]}.art.{key}',art.get(key),row[col],f'Decor Art Briefs!{openpyxl.utils.get_column_letter(col+1)}{number}')
        same(f'{row[0]}.static_policy',art.get('motion',{}).get('kind')=='static',row[8]=='Static',f'Decor Art Briefs!I{number}')
    same('events',len(data['events']),12,'Events!A5:A16')
    for event,row in zip(data['events'],sheets['Events'][4:16]):
        for key,col in {'name':1,'proposed_window':2,'coin_id':3,'pearl_id':4}.items():same('event.'+key,event.get(key),row[col],'Events')
        same('event.default_closed',event.get('configured'),False,'Events require an explicit release')
    same('no extra tutorial purchase XP',data['tuning']['tutorial_total_xp'],0,'Normal first-purchase reward only')
    result=dict(checks=checks,errors=errors,items=len(items),events=len(data['events']),workbook=str(args.workbook))
    args.out.parent.mkdir(parents=True,exist_ok=True);args.out.write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2));return bool(errors)

if __name__=='__main__':raise SystemExit(main())
